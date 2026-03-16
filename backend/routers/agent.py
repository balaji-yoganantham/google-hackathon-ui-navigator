"""Agent API: execute, status, stream SSE, WebSocket live feed, continue, cancel, history."""
import asyncio
import json
import time
import uuid
from datetime import datetime
from typing import Any, Optional

from fastapi import APIRouter, HTTPException, WebSocket, WebSocketDisconnect
from fastapi.requests import Request
from fastapi.responses import Response, StreamingResponse
from pydantic import BaseModel

from agents.orchestrator import run_task
from models.schemas import ExtractionResult, TaskExecution
from utils.browser_pool import browser_pool
from utils.session_manager import session_manager
from utils.stream_manager import browser_stream_manager
from utils.task_queue import TaskQueue

router = APIRouter(prefix="/api/agent", tags=["agent"])

# Cancelled session IDs (cleared when task ends)
_cancelled_sessions: set[str] = set()


def _make_task_dict(task: TaskExecution) -> dict:
    return task.model_dump(mode="json")


async def _queue_processor(session_id: str, payload: Any) -> None:
    kind = payload.get("kind")
    data = payload.get("data") or {}
    if kind == "execute":
        task_description = data.get("taskDescription", "")
        start_url = data.get("startUrl", "") or ""
        if not task_description:
            session_manager.update_session(
                session_id,
                TaskExecution(
                    id=session_id,
                    taskDescription=task_description or "",
                    status="failed",
                    steps=[],
                    currentScreenshot="",
                    error="taskDescription is required",
                    startUrl=start_url,
                    createdAt=datetime.utcnow(),
                    updatedAt=datetime.utcnow(),
                ),
            )
            return

        def on_update(t: TaskExecution) -> None:
            session_manager.update_session(session_id, t)

        def get_cancelled() -> bool:
            return session_id in _cancelled_sessions

        try:
            task = await run_task(
                session_id,
                task_description,
                start_url,
                on_update=on_update,
                get_cancelled=get_cancelled,
            )
            session_manager.update_session(session_id, task)
        except ValueError as e:
            session_manager.update_session(
                session_id,
                TaskExecution(
                    id=session_id,
                    taskDescription=task_description,
                    status="failed",
                    steps=[],
                    currentScreenshot="",
                    error=str(e),
                    startUrl=start_url,
                    createdAt=datetime.utcnow(),
                    updatedAt=datetime.utcnow(),
                ),
            )
        except json.JSONDecodeError as e:
            session_manager.update_session(
                session_id,
                TaskExecution(
                    id=session_id,
                    taskDescription=task_description,
                    status="failed",
                    steps=[],
                    currentScreenshot="",
                    error="Model response was not valid. Try a shorter or simpler task.",
                    startUrl=start_url,
                    createdAt=datetime.utcnow(),
                    updatedAt=datetime.utcnow(),
                ),
            )
        finally:
            _cancelled_sessions.discard(session_id)
    elif kind == "continue":
        entry = session_manager.get_session_entry(session_id)
        if not entry:
            session_manager.update_session(
                session_id,
                TaskExecution(
                    id=session_id,
                    taskDescription=data.get("instruction", ""),
                    status="failed",
                    steps=[],
                    currentScreenshot="",
                    error="Session not found for continue",
                    createdAt=datetime.utcnow(),
                    updatedAt=datetime.utcnow(),
                ),
            )
            return
        existing_task, _ = entry
        instruction = data.get("instruction", "")
        start_url = existing_task.startUrl or ""
        if not start_url:
            start_url = "about:blank"
        if not instruction:
            session_manager.update_session(
                session_id,
                TaskExecution(
                    id=session_id,
                    taskDescription=instruction,
                    status="failed",
                    steps=[],
                    currentScreenshot=existing_task.currentScreenshot,
                    error="instruction is required",
                    startUrl=start_url,
                    createdAt=existing_task.createdAt,
                    updatedAt=datetime.utcnow(),
                ),
            )
            return

        def on_update(t: TaskExecution) -> None:
            session_manager.update_session(session_id, t)

        def get_cancelled() -> bool:
            return session_id in _cancelled_sessions

        task = await run_task(
            session_id,
            instruction,
            start_url,
            on_update=on_update,
            get_cancelled=get_cancelled,
        )
        session_manager.update_session(session_id, task)
        _cancelled_sessions.discard(session_id)
    else:
        session_manager.update_session(
            session_id,
            TaskExecution(
                id=session_id,
                taskDescription="",
                status="failed",
                steps=[],
                currentScreenshot="",
                error="Unknown payload kind",
                createdAt=datetime.utcnow(),
                updatedAt=datetime.utcnow(),
            ),
        )


queue = TaskQueue(_queue_processor)


class ExecuteBody(BaseModel):
    taskDescription: str
    startUrl: Optional[str] = None


class ContinueBody(BaseModel):
    instruction: str
    inputs: dict | None = None


@router.post("/execute")
async def execute_task(body: ExecuteBody) -> dict:
    if not body.taskDescription:
        raise HTTPException(400, "taskDescription is required")
    await browser_pool.initialize()
    session_id = uuid.uuid4().hex[:8]
    pending = TaskExecution(
        id=session_id,
        taskDescription=body.taskDescription,
        status="pending",
        steps=[],
        currentScreenshot="",
        startUrl=body.startUrl or "",
        createdAt=datetime.utcnow(),
        updatedAt=datetime.utcnow(),
    )
    session_manager.create_session(session_id, pending, None)
    queue.enqueue(
        session_id,
        {"kind": "execute", "data": {"taskDescription": body.taskDescription, "startUrl": body.startUrl or "", "uid": None}},
    )
    return {
        "sessionId": session_id,
        "task": _make_task_dict(pending),
        "message": "Task queued for execution",
        "queueLength": queue.get_queue_length(),
    }


@router.get("/status/{session_id}")
async def get_task_status(session_id: str) -> dict:
    """GET /api/agent/status/{session_id}. Returns 404 if session is missing (e.g. created on another instance)."""
    task = session_manager.get_session(session_id)
    if not task:
        raise HTTPException(404, "Session not found")
    return _make_task_dict(task)


# Origins allowed for CORS on SSE (must match main.py allow_origins for credentials)
_STREAM_CORS_ORIGINS = frozenset({
    "http://localhost:5173",
    "http://127.0.0.1:5173",
    "http://localhost:8080",
    "http://127.0.0.1:8080",
    "https://project-70591921-7d7a-4043-ba8.web.app",
    "https://project-70591921-7d7a-4043-ba8.firebaseapp.com",
})


@router.get("/stream/{session_id}")
async def stream_task(session_id: str, request: Request) -> StreamingResponse:
    async def event_stream():
        last_steps = -1
        last_status = ""
        last_reports_len = -1
        last_heartbeat = time.monotonic()
        heartbeat_interval = 15.0  # send keepalive so proxies don't close the stream
        while True:
            if await request.is_disconnected():
                break
            task = session_manager.get_session(session_id)
            if not task:
                yield f"data: {json.dumps({'error': 'Session not found'})}\n\n"
                break
            steps_len = len(task.steps)
            reports_len = len(task.reports) if task.reports else 0
            if steps_len != last_steps or task.status != last_status or reports_len != last_reports_len:
                last_steps = steps_len
                last_status = task.status
                last_reports_len = reports_len
                yield f"data: {json.dumps(_make_task_dict(task))}\n\n"
            if task.status in ("completed", "failed", "cancelled"):
                break
            now = time.monotonic()
            if now - last_heartbeat >= heartbeat_interval:
                yield ": keepalive\n\n"
                last_heartbeat = now
            await asyncio.sleep(0.4)

    origin = request.headers.get("origin", "")
    cors_headers = {
        "Cache-Control": "no-cache",
        "Connection": "keep-alive",
        "X-Accel-Buffering": "no",
    }
    if origin in _STREAM_CORS_ORIGINS:
        cors_headers["Access-Control-Allow-Origin"] = origin
        cors_headers["Access-Control-Allow-Credentials"] = "true"

    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers=cors_headers,
    )


@router.websocket("/ws/{session_id}")
async def browser_stream_ws(websocket: WebSocket, session_id: str) -> None:
    """
    WebSocket live browser feed.

    Clients connect here to receive real-time JPEG frames as the agent acts.
    Each message is JSON: {"type": "screenshot", "screenshot": "<base64>", "step": N}
    Send any message to keep the connection alive; send "stop" to close it.
    """
    await websocket.accept()
    browser_stream_manager.connect(session_id, websocket)
    try:
        while True:
            data = await websocket.receive_text()
            if data.strip().lower() == "stop":
                break
    except WebSocketDisconnect:
        pass
    finally:
        browser_stream_manager.disconnect(session_id)


@router.post("/continue/{session_id}")
async def continue_task(session_id: str, body: ContinueBody) -> dict:
    if not body.instruction:
        raise HTTPException(400, "instruction is required")
    entry = session_manager.get_session_entry(session_id)
    if not entry:
        raise HTTPException(404, "Session not found")
    existing_task, _ = entry
    if existing_task.status in ("running", "pending"):
        raise HTTPException(409, "Task is already running")
    await browser_pool.initialize()
    pending = TaskExecution(
        id=existing_task.id,
        taskDescription=body.instruction,
        status="pending",
        steps=existing_task.steps,
        currentScreenshot=existing_task.currentScreenshot,
        startUrl=existing_task.startUrl,
        createdAt=existing_task.createdAt,
        updatedAt=datetime.utcnow(),
    )
    session_manager.update_session(session_id, pending)
    queue.enqueue(
        session_id,
        {"kind": "continue", "data": {"instruction": body.instruction, "inputs": body.inputs, "uid": None}},
    )
    return {
        "sessionId": session_id,
        "task": _make_task_dict(pending),
        "message": "Continue queued for execution",
        "queueLength": queue.get_queue_length(),
    }


@router.post("/cancel/{session_id}")
async def cancel_task(session_id: str) -> dict:
    entry = session_manager.get_session_entry(session_id)
    if not entry:
        raise HTTPException(404, "Session not found")
    task, _ = entry
    _cancelled_sessions.add(session_id)
    removed_from_queue = queue.cancel(session_id)
    is_running = queue.get_current_session_id() == session_id
    cancelled_task = TaskExecution(
        id=task.id,
        taskDescription=task.taskDescription,
        status="cancelled",
        steps=task.steps,
        currentScreenshot=task.currentScreenshot,
        startUrl=task.startUrl,
        createdAt=task.createdAt,
        updatedAt=datetime.utcnow(),
    )
    session_manager.update_session(session_id, cancelled_task)
    if is_running:
        await browser_pool.close()
        await browser_pool.initialize()
    return {
        "message": "Task cancelled",
        "sessionId": session_id,
        "removedFromQueue": removed_from_queue,
        "wasRunning": is_running,
        "task": _make_task_dict(cancelled_task),
    }


@router.post("/cleanup/{session_id}")
async def cleanup_session(session_id: str) -> dict:
    session_manager.delete_session(session_id)
    return {"message": "Session cleaned up"}


@router.get("/history")
async def get_task_history() -> list:
    """Return list of session IDs (frontend can fetch each via status or history/:id)."""
    return list(session_manager.list_session_ids())


@router.get("/history/{task_id}")
async def get_task_detail(task_id: str) -> dict:
    task = session_manager.get_session(task_id)
    if not task:
        raise HTTPException(404, "Task not found")
    return _make_task_dict(task)


@router.delete("/history/{task_id}")
async def delete_task_history(task_id: str) -> dict:
    session_manager.delete_session(task_id)
    return {"message": "Deleted"}


@router.get("/report/{session_id}", response_model=ExtractionResult)
async def get_report(session_id: str) -> ExtractionResult:
    """Return extraction reports for a session (from completed extract step)."""
    task = session_manager.get_session(session_id)
    if not task:
        raise HTTPException(404, "Session not found")
    return ExtractionResult(reports=task.reports, session_id=session_id)


def _report_type_subtitle(reports: list) -> str:
    """Return report type label from content_type and count (e.g. Legal Research Report, Video Analysis Report)."""
    if not reports:
        return "Extraction Report"
    ct = (reports[0].content_type or "page").lower()
    n = len(reports)
    if ct == "pdf":
        return "Legal Research Report" if n <= 1 else f"Precedent Research Report — {n} Cases"
    if ct == "youtube":
        return "Video Analysis Report"
    if ct == "qa":
        return "QA Scan Report" if n <= 1 else f"QA Scan Report — {n} issues"
    return "Legal Research Report" if n <= 1 else f"Extracted Cases ({n})"


@router.post("/export/{session_id}")
async def export_docx(session_id: str) -> Response:
    """Generate and return a .docx file from the session's extraction reports (title block, Research Query, court/date/docket)."""
    task = session_manager.get_session(session_id)
    if not task:
        raise HTTPException(404, "Session not found")
    if not task.reports:
        raise HTTPException(404, "No reports to export")
    import io
    from datetime import datetime as dt
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading("UI Navigator", 0)
    doc.add_paragraph(_report_type_subtitle(task.reports))
    doc.add_paragraph(f"Generated {dt.utcnow().strftime('%Y-%m-%d %H:%M UTC')}")
    doc.add_paragraph()
    if getattr(task, "taskDescription", None) and (task.taskDescription or "").strip():
        doc.add_heading("Research Query", level=1)
        doc.add_paragraph((task.taskDescription or "").strip())
        doc.add_paragraph()
    for i, r in enumerate(task.reports, 1):
        doc.add_heading(r.title or f"Report {i}", level=1)
        meta_parts = []
        if r.court:
            meta_parts.append(f"Court: {r.court}")
        if r.date:
            meta_parts.append(f"Date: {r.date}")
        if r.docket:
            meta_parts.append(f"Docket: {r.docket}")
        if meta_parts:
            doc.add_paragraph(" | ".join(meta_parts))
        if r.url and not any(r.url in p for p in meta_parts):
            doc.add_paragraph(f"URL: {r.url}")
        for block in (r.content or "").split("\n"):
            block = block.strip()
            if block:
                doc.add_paragraph(block)
        doc.add_paragraph()
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="report-{session_id}.docx"'},
    )
